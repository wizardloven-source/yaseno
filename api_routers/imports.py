# api_routers/imports.py
"""
Import Center Router - مركز استيراد Excel

نقطة واحدة لرفع ملفات Excel واستيراد البيانات الرئيسية دفعة واحدة:
العملاء، الموردون، المنتجات، دليل الحسابات، الفروع/المواقع، مراكز التكلفة،
المشاريع، العملات.

السلوك الديناميكي:
- رؤوس الأعمدة تُكتشف تلقائياً (أو تُمرَّر خريطة المطابقة من الواجهة).
- الأكواد الناقصة/القصيرة تُولَّد تلقائياً من المستودعات.
- الأسماء الناقصة تُولَّد تلقائياً.
- تفاصيل غير موجودة تأخذ قيماً افتراضية (العملة الأساسية، الدولة...).
- كل صف مستقل: فشل صف لا يوقف بقية الصفوف، وتُجمَع النتائج لكل صف.
"""

import io
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Callable, Dict, List, Optional

from fastapi import APIRouter, Depends, File, Form, UploadFile

from api_routers.shared import bootstrap, logger, ApiResponse, get_current_user
from api_routers.shared.auth_deps import require_permission

router = APIRouter(prefix="", tags=["import"])

BASE_CURRENCY_FALLBACK = "USD"
SUPPORTED_ENTITIES = {
    "customers",
    "suppliers",
    "products",
    "accounts",
    "sites",
    "centers",
    "projects",
    "currencies",
}

# =============================================================================
# أدوات خلايا Excel
# =============================================================================


def _cell_to_str(cell: Any) -> str:
    """تحويل أي خلية (None/رقم/تاريخ/نص) إلى نص نظيف."""
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d")
    if hasattr(cell, "isoformat"):
        try:
            return cell.isoformat()
        except Exception:
            return str(cell)
    if isinstance(cell, bool):
        return "1" if cell else "0"
    if isinstance(cell, Decimal):
        return f"{cell.normalize():f}" if cell == cell.to_integral() else str(cell)
    if isinstance(cell, float):
        if cell.is_integer():
            return str(int(cell))
        return str(cell)
    return str(cell)


def _clean(v: Optional[str]) -> Optional[str]:
    """إزالة المسافات؛ تُرجع None للفارغ."""
    if v is None:
        return None
    t = v.strip()
    return t if t else None


def _parse_number(v: Optional[str]) -> Optional[Decimal]:
    if not v:
        return None
    s = v.strip().replace(",", "").replace("٫", ".")
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _parse_int(v: Optional[str]) -> Optional[int]:
    n = _parse_number(v)
    if n is None:
        return None
    try:
        return int(n)
    except (ValueError, TypeError):
        return None


def _parse_bool(v: Optional[str]) -> Optional[bool]:
    if not v:
        return None
    s = v.strip().lower()
    if s in ("1", "true", "yes", "y", "نعم", "مفعل", "active"):
        return True
    if s in ("0", "false", "no", "n", "لا", "معطل", "inactive"):
        return False
    return None


def _parse_date(v: Optional[str]) -> Optional[str]:
    """يحوّل نص تاريخ إلى YYYY-MM-DD عبر عدة صيغ، أو None."""
    if not v:
        return None
    s = v.strip().replace("/", "-").replace(".", "-").split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y"):
        try:
            d = datetime.strptime(s, fmt)
            return d.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _norm_header(h: str) -> str:
    return h.strip().lower().replace("-", "").replace("_", "").replace(" ", "")


# =============================================================================
# أسماء حقول لكل كيان (للتحقق والمطابقة التلقائية عند غياب الـ mapping)
# =============================================================================

ENTITY_ALIASES: Dict[str, Dict[str, List[str]]] = {
    "customers": {
        "code": ["code", "customer_code", "customer code", "الرمز", "الكود", "رقم العميل", "كود العميل"],
        "name": ["name", "customer_name", "customer name", "الاسم", "اسم العميل"],
        "phone": ["phone", "tel", "telephone", "الهاتف", "رقم الهاتف"],
        "mobile": ["mobile", "mobile_no", "الجوال", "الموبايل"],
        "email": ["email", "e-mail", "mail", "البريد", "البريد الالكتروني", "ايميل"],
        "street": ["street", "address", "العنوان", "الشارع"],
        "city": ["city", "المدينة", "المحافظة"],
        "country": ["country", "الدولة", "البلد"],
        "tax_number": ["tax_number", "tax", "vat", "الرقم الضريبي"],
        "credit_limit": ["credit_limit", "credit limit", "حد الائتمان", "سقف الائتمان"],
        "currency": ["currency", "العملة", "عملة"],
        "branches": ["branches", "branch", "branch names", "الفروع", "أسماء الفروع", "اسم الفرع", "الأفرع", "فروع العميل"],
        "notes": ["notes", "note", "ملاحظات", "بيان"],
    },
    "suppliers": {
        "code": ["code", "supplier_code", "supplier code", "الرمز", "الكود", "كود المورد"],
        "name": ["name", "supplier_name", "supplier name", "الاسم", "اسم المورد"],
        "phone": ["phone", "tel", "telephone", "الهاتف", "رقم الهاتف"],
        "mobile": ["mobile", "mobile_no", "الجوال", "الموبايل"],
        "email": ["email", "e-mail", "mail", "البريد", "البريد الالكتروني", "ايميل"],
        "street": ["street", "address", "العنوان", "الشارع"],
        "city": ["city", "المدينة", "المحافظة"],
        "country": ["country", "الدولة", "البلد"],
        "tax_number": ["tax_number", "tax", "vat", "الرقم الضريبي"],
        "credit_limit": ["credit_limit", "credit limit", "حد الائتمان", "سقف الائتمان"],
        "currency": ["currency", "العملة", "عملة"],
        "notes": ["notes", "note", "ملاحظات", "بيان"],
    },
    "products": {
        "code": ["code", "product_code", "product code", "sku", "الرمز", "الكود", "كود المنتج"],
        "name": ["name", "product_name", "product name", "الاسم", "اسم المنتج"],
        "unit_price": ["unit_price", "unit price", "price", "السعر", "سعر الوحدة"],
        "tax_rate": ["tax_rate", "tax rate", "tax", "vat", "الضريبة", "نسبة الضريبة"],
        "description": ["description", "details", "الوصف"],
        "category": ["category", "categories", "type", "التصنيف", "الفئة"],
        "stock_quantity": ["stock_quantity", "stock", "quantity", "qty", "الكمية", "المخزون"],
        "low_stock_threshold": ["low_stock_threshold", "low stock", "threshold", "حد التنبيه"],
        "currency": ["currency", "العملة", "عملة"],
    },
    "accounts": {
        "code": ["code", "account_code", "account code", "الكود", "الرمز", "كود الحساب"],
        "name": ["name", "account_name", "account name", "الاسم", "اسم الحساب"],
        "account_type": ["account_type", "account type", "type", "نوع الحساب"],
        "parent_code": ["parent_code", "parent", "parent account", "الحساب الرئيسي", "الحساب الأب"],
        "description": ["description", "details", "الوصف"],
        "currency": ["currency", "العملة", "عملة"],
        "is_active": ["is_active", "active", "مفعل"],
    },
    "sites": {
        "code": ["code", "site_code", "site code", "الكود", "الرمز", "كود الموقع"],
        "name": ["name", "site_name", "site name", "الاسم", "اسم الموقع"],
        "site_type": ["site_type", "site type", "type", "نوع الموقع"],
        "street": ["street", "address", "العنوان", "الشارع"],
        "city": ["city", "المدينة", "المحافظة"],
        "country": ["country", "الدولة", "البلد"],
        "phone": ["phone", "tel", "الهاتف"],
        "mobile": ["mobile", "الجوال", "الموبايل"],
        "email": ["email", "e-mail", "mail", "البريد", "ايميل"],
        "contact_person": ["contact_person", "contact person", "manager_name", "manager name", "مدير الموقع", "المسؤول", "شخص الاتصال"],
        "notes": ["notes", "note", "ملاحظات"],
        "is_default": ["is_default", "default", "الافتراضي", "الرئيسي"],
    },
    "centers": {
        "code": ["code", "center_code", "center code", "الكود", "الرمز", "كود المركز"],
        "name": ["name", "center_name", "center name", "الاسم", "اسم المركز"],
        "center_type": ["center_type", "center type", "type", "نوع المركز"],
        "parent_code": ["parent_code", "parent", "parent center", "المركز الرئيسي", "المركز الأب"],
        "manager_name": ["manager_name", "manager", "مدير المركز"],
        "department": ["department", "القسم", "الإدارة"],
        "budget_amount": ["budget_amount", "budget", "الميزانية"],
        "budget_currency": ["budget_currency", "budget currency", "عملة الميزانية"],
        "description": ["description", "details", "الوصف"],
    },
    "projects": {
        "code": ["code", "project_code", "project code", "الكود", "الرمز", "كود المشروع"],
        "name": ["name", "project_name", "project name", "الاسم", "اسم المشروع"],
        "description": ["description", "details", "الوصف"],
        "customer_name": ["customer_name", "customer", "العميل", "اسم العميل"],
        "manager_name": ["manager_name", "manager", "مدير المشروع"],
        "budget_amount": ["budget_amount", "budget", "الميزانية"],
        "budget_currency": ["budget_currency", "budget currency", "عملة الميزانية"],
        "start_date": ["start_date", "start date", "تاريخ البدء", "البداية"],
        "end_date": ["end_date", "end date", "تاريخ الانتهاء", "النهاية"],
        "status": ["status", "الحالة", "الحالة"],
        "tags": ["tags", "tag", "الوسوم"],
    },
    "currencies": {
        "code": ["code", "currency_code", "currency code", "الكود", "الرمز", "كود العملة"],
        "name": ["name", "currency_name", "currency name", "الاسم", "اسم العملة"],
        "symbol": ["symbol", "الرمز", "العملة"],
        "decimal_places": ["decimal_places", "decimal", "المنازل العشرية"],
        "is_base": ["is_base", "base", "الأساسية", "الرئيسية"],
    },
}


def _auto_map(headers: List[str], entity: str, mapping: Dict[str, int]) -> Dict[str, int]:
    """يملأ حقولاً غير معيَّنة من المطابقة التلقائية برؤوس الأعمدة."""
    aliases = ENTITY_ALIASES.get(entity, {})
    normalized = {k: [_norm_header(a) for a in vals] for k, vals in aliases.items()}
    for field, alist in normalized.items():
        if field in mapping and mapping[field] >= 0:
            continue
        for i, h in enumerate(headers):
            if _norm_header(h) in alist:
                mapping[field] = i
                break
    return mapping


# =============================================================================
# مولّدات الأكواد
# =============================================================================


def _safe_next_code(uow: Any, repo_attr: str, prefix: str) -> str:
    """يولّد الكود التالي بأمان عبر get_next_code أو رأسياً يدوياً."""
    repo = getattr(uow, repo_attr, None)
    try:
        if repo and hasattr(repo, "get_next_code"):
            candidate = repo.get_next_code(prefix=prefix)
            if candidate:
                return candidate
    except Exception as e:
        logger.warning(f"get_next_code failed ({repo_attr}): {e}")
    # يدوي: بادئة + مقطع زمني بسيط لضمان التفرد
    return f"{prefix}{int(datetime.now().timestamp() * 1000) % 100000:05d}"


def _next_account_code(uow: Any, account_type: str) -> str:
    """توليد كود حساب رقمي حسب نوع الحساب (1/2/3/4/5 + تسلسل)."""
    prefix_digit = {"asset": "1", "liability": "2", "equity": "3", "revenue": "4", "expense": "5"}.get(account_type or "", "1")
    max_num = 0
    try:
        from sqlalchemy import func, select
        from core.infrastructure.db.models.account_model import AccountModel

        row = uow.session.execute(
            select(func.max(AccountModel.code)).where(
                AccountModel.code.like(f"{prefix_digit}%") &
                AccountModel.code.op("~")("^[0-9]+$")
            )
        ).scalar_one_or_none()
        if row:
            try:
                max_num = int(str(row))
            except ValueError:
                max_num = 0
    except Exception as e:
        logger.warning(f"account code scan failed: {e}")
    seq = max_num + 1 if max_num >= 1000 else 1
    return f"{prefix_digit}{seq:04d}"


# =============================================================================
# إنشاء الصفوف (نمط: إما uow مباشر أو command_bus مثل الراوترات)
# =============================================================================


class ImportRowError(Exception):
    pass


def _split_branches(raw: Optional[str]) -> List[str]:
    """يقسّم نص الفروع إلى أسماء مفصولة بأيٍ من: الفاصلة، النقطة، الفاصلة المنقوطة،
    الشرطة، الشرطة السفلية، أو سطر جديد (مع إزالة الفراغ وتجاهل الأسماء القصيرة/المكررة)."""
    if not raw:
        return []
    parts = re.split(r"[,;،.\-_\n\r\t]+", raw)
    seen: set[str] = set()
    out: List[str] = []
    for p in parts:
        name = p.strip()
        if len(name) >= 2 and name not in seen:
            seen.add(name)
            out.append(name)
    return out


def _create_customer_branches(uow: Any, customer: Any, raw_branches: Optional[str], username: str) -> int:
    """ينشئ فروعاً مرتبطة بالعميل من عمود 'الفروع'؛ يُرجع عدد الفروع المنشأة.

    الفروع تُربط بمعرف العميل نفسه (customer_id) فتظهر في نافذة الفروع مباشرة.
    الأسماء المكررة للعميل نفسه تُتجاهل، والأول يُعيّن افتراضياً إن لم يوجد افتراضي سابق.
    """
    names = _split_branches(raw_branches)
    if not names:
        return 0
    repo = getattr(uow, "customer_branches", None)
    if repo is None:
        return 0
    from core.domain.customer_branch.entities import CustomerBranch

    customer_id = str(customer.id.value) if hasattr(customer.id, "value") else str(customer.id)
    customer_name = getattr(customer, "name", None) or ""
    customer_code = getattr(customer.code, "value", None) or ""

    existing_names = {
        (b.name or "").strip()
        for b in repo.get_by_customer(customer_id, include_inactive=True, limit=1000)
    }
    has_default = repo.get_default_branch(customer_id) is not None

    created = 0
    for name in names:
        if name in existing_names:
            continue
        code = repo.get_next_code(prefix="BR")
        branch = CustomerBranch.create(
            code=code,
            name=name,
            customer_id=customer_id,
            customer_name=customer_name,
            customer_code=customer_code,
            branch_type="store",
            is_default=(not has_default and created == 0),
            created_by=username,
        )
        repo.save(branch)
        created += 1
    return created


def _create_customer(uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.domain.customers.entities import Customer
    from core.domain.customers.value_objects import CustomerCode, ContactInfo, Address

    code = _clean(values.get("code"))
    if code is None or len(code) < 3 or len(code) > 20:
        code = _safe_next_code(uow, "customers", "C")
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"عميل {code}"

    existing = uow.customers.get_by_code(CustomerCode(code)) if hasattr(uow.customers, "get_by_code") else None
    if existing:
        if not options.get("update_existing", True):
            return f"تم تجاهل العميل '{code}' (موجود مسبقاً)"
        updates = {}
        n = _clean(values.get("name"))
        if n:
            updates["name"] = n
        em = _clean(values.get("email"))
        if em:
            updates["email"] = em
        ph = _clean(values.get("phone"))
        if ph:
            updates["phone"] = ph
        if updates:
            if "name" in updates:
                existing.name = updates["name"]
            contact = existing.contact_info
            existing.contact_info = ContactInfo(
                email=updates.get("email", contact.email if contact else None),
                phone=updates.get("phone", contact.phone if contact else None),
                mobile=contact.mobile if contact else None,
            )
            uow.customers.save(existing)
            created_branches = _create_customer_branches(uow, existing, values.get("branches"), username)
            uow.commit()
            suffix = f" مع {created_branches} فروع" if created_branches else ""
            return f"تم تحديث العميل '{code}'{suffix}"

    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    cc = _clean(values.get("currency"))
    currency = cc if cc and len(cc) == 3 else base_currency
    country = _clean(values.get("country")) or "LB"
    cl = _parse_number(values.get("credit_limit"))
    customer = Customer.create(
        code=CustomerCode(code),
        name=name,
        contact_info=ContactInfo(
            email=_clean(values.get("email")),
            phone=_clean(values.get("phone")),
            mobile=_clean(values.get("mobile")),
        ),
        address=Address(
            street=_clean(values.get("street")),
            city=_clean(values.get("city")),
            country=country,
        ),
        tax_number=_clean(values.get("tax_number")),
        credit_limit=Decimal(cl if cl is not None else 0),
        currency=currency,
        notes=_clean(values.get("notes")),
        created_by=username,
    )
    uow.customers.save(customer)
    created_branches = _create_customer_branches(uow, customer, values.get("branches"), username)
    uow.commit()
    suffix = f" مع {created_branches} فروع" if created_branches else ""
    return f"تم إنشاء العميل '{code}'{suffix}"


def _create_supplier(uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.domain.suppliers.entities import Supplier
    from core.domain.suppliers.value_objects import SupplierCode, ContactInfo, Address

    code = _clean(values.get("code"))
    if code is None or len(code) < 1 or len(code) > 20:
        code = _safe_next_code(uow, "suppliers", "S")
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"مورد {code}"

    existing = uow.suppliers.get_by_code(SupplierCode(code)) if hasattr(uow.suppliers, "get_by_code") else None
    if existing:
        return f"تم تجاهل المورد '{code}' (موجود مسبقاً)"

    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    cc = _clean(values.get("currency"))
    currency = cc if cc and len(cc) == 3 else base_currency
    country = _clean(values.get("country")) or "LB"
    cl = _parse_number(values.get("credit_limit"))
    supplier = Supplier.create(
        code=SupplierCode(code),
        name=name,
        contact_info=ContactInfo(
            email=_clean(values.get("email")),
            phone=_clean(values.get("phone")),
            mobile=_clean(values.get("mobile")),
        ),
        address=Address(
            street=_clean(values.get("street")),
            city=_clean(values.get("city")),
            country=country,
        ),
        tax_number=_clean(values.get("tax_number")),
        credit_limit=Decimal(cl if cl is not None else 0),
        currency=currency,
        notes=_clean(values.get("notes")),
        created_by=username,
    )
    uow.suppliers.save(supplier)
    uow.commit()
    return f"تم إنشاء المورد '{code}'"


def _create_product(uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.domain.products.entities import Product
    from core.domain.products.value_objects import ProductCode
    from core.domain.shared.value_objects import Money

    code = _clean(values.get("code"))
    if code is None or len(code) < 1:
        code = _safe_next_code(uow, "products", "P")
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"منتج {code}"

    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    cc = _clean(values.get("currency"))
    currency = cc if cc and len(cc) == 3 else base_currency
    price = _parse_number(values.get("unit_price"))
    tax_rate = _parse_number(values.get("tax_rate"))
    stock = _parse_int(values.get("stock_quantity"))
    low = _parse_int(values.get("low_stock_threshold"))

    existing = uow.products.get_by_code(ProductCode(code)) if hasattr(uow.products, "get_by_code") else None
    if existing:
        same_name = (existing.name or "").strip().lower() == (name or "").strip().lower()
        if not same_name:
            code = _safe_next_code(uow, "products", "P")
        else:
            if not options.get("update_existing", True):
                return f"تم تجاهل المنتج '{code}' (موجود مسبقاً)"
            if price is not None:
                existing.unit_price = Money(price, currency)
            if tax_rate is not None:
                existing.tax_rate = tax_rate
            if stock is not None:
                existing.stock_quantity = stock
            if low is not None:
                existing.low_stock_threshold = low
            desc = _clean(values.get("description"))
            if desc is not None:
                existing.description = desc
            cat = _clean(values.get("category"))
            if cat is not None:
                existing.category = cat
            n = _clean(values.get("name"))
            if n and len(n) >= 2:
                existing.name = n
            uow.products.save(existing)
            uow.commit()
            return f"تم تحديث المنتج '{code}'"

    product = Product.create(
        code=ProductCode(code),
        name=name,
        unit_price=Money(Decimal(price if price is not None else 0), currency),
        tax_rate=Decimal(tax_rate if tax_rate is not None else 0),
        description=_clean(values.get("description")),
        category=_clean(values.get("category")),
        stock_quantity=stock if stock is not None else 0,
        low_stock_threshold=low if low is not None else 10,
        created_by=username,
    )
    uow.products.save(product)
    uow.commit()
    return f"تم إنشاء المنتج '{code}'"


def _dispatch_command(command_bus: Any, command_class, **kwargs):
    command = command_class(**kwargs)
    return command_bus.dispatch(command)


def _create_account( command_bus: Any, uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.application.accounts.commands import CreateAccountCommand

    account_type = (_clean(values.get("account_type")) or "asset").lower()
    mapping_t = {
        "assets": "asset", "اصل": "asset", "اصول": "asset", "أصول": "asset",
        "liabilities": "liability", "خصوم": "liability", "مطلوبات": "liability", "خصم": "liability",
        "equity": "equity", "حقق": "equity", "حقوق": "equity", "حقوق الملكية": "equity", "ملكية": "equity",
        "revenue": "revenue", "ايرادات": "revenue", "إيرادات": "revenue", "دخل": "revenue", "income": "revenue",
        "expense": "expense", "مصاريف": "expense", "مصروفات": "expense",
    }
    at = mapping_t.get(account_type, account_type)
    if at not in ("asset", "liability", "equity", "revenue", "expense"):
        at = "asset"

    code = _clean(values.get("code"))
    if code is None or not re.match(r"^\d{3,20}$", code):
        code = _next_account_code(uow, at)
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"حساب {code}"

    existing = None
    try:
        from core.domain.shared.value_objects import AccountCode
        existing = uow.accounts.get_by_code(AccountCode(code))
    except Exception as e:
        logger.debug(f"account dup check failed: {e}")
    if existing:
        return f"تم تجاهل الحساب '{code}' (موجود مسبقاً)"

    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    cc = _clean(values.get("currency"))
    currency = cc if cc and len(cc) == 3 else base_currency
    is_active = _parse_bool(values.get("is_active"))
    _dispatch_command(
        command_bus, CreateAccountCommand,
        code=code, name=name, account_type=at,
        parent_code=_clean(values.get("parent_code")),
        description=_clean(values.get("description")),
        currency=currency, is_active=is_active if is_active is not None else True,
        created_by=username,
    )
    return f"تم إنشاء الحساب '{code}'"


def _create_site(command_bus: Any, uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.application.sites.commands import CreateSiteCommand

    code = _clean(values.get("code"))
    if code is None or len(code) < 1:
        code = _safe_next_code(uow, "sites", "S")
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"موقع {code}"

    existing = None
    try:
        from core.domain.sites.value_objects import SiteCode
        existing = uow.sites.get_by_code(SiteCode(code))
    except Exception as e:
        logger.debug(f"site dup check failed: {e}")
    if existing:
        return f"تم تجاهل الموقع '{code}' (موجود مسبقاً)"

    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    cc = _clean(values.get("currency"))
    currency = cc if cc and len(cc) == 3 else base_currency
    is_default = _parse_bool(values.get("is_default"))
    _dispatch_command(
        command_bus, CreateSiteCommand,
        code=code, name=name,
        site_type=_clean(values.get("site_type")) or "general",
        street=_clean(values.get("street")),
        city=_clean(values.get("city")),
        country=_clean(values.get("country")) or "LB",
        phone=_clean(values.get("phone")),
        mobile=_clean(values.get("mobile")),
        email=_clean(values.get("email")),
        contact_person=_clean(values.get("contact_person")),
        notes=_clean(values.get("notes")),
        is_default=is_default if is_default is not None else False,
        created_by=username,
    )
    return f"تم إنشاء الموقع '{code}'"


def _create_center(command_bus: Any, uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.application.centers.commands import CreateCenterCommand

    code = _clean(values.get("code"))
    if code is None or len(code) < 1:
        code = _safe_next_code(uow, "centers", "C")
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = f"مركز {code}"

    existing = None
    try:
        from core.domain.centers.value_objects import CenterCode
        existing = uow.centers.get_by_code(CenterCode(code))
    except Exception as e:
        logger.debug(f"center dup check failed: {e}")
    if existing:
        return f"تم تجاهل المركز '{code}' (موجود مسبقاً)"

    center_type = (_clean(values.get("center_type")) or "cost").lower()
    if center_type not in ("cost", "profit", "both"):
        center_type = "cost"
    budget = _parse_number(values.get("budget_amount"))
    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    bcc = _clean(values.get("budget_currency"))
    budget_currency = bcc if bcc and len(bcc) == 3 else base_currency
    _dispatch_command(
        command_bus, CreateCenterCommand,
        code=code, name=name, center_type=center_type,
        parent_code=_clean(values.get("parent_code")),
        manager_name=_clean(values.get("manager_name")),
        department=_clean(values.get("department")),
        budget_amount=Decimal(budget if budget is not None else 0),
        budget_currency=budget_currency,
        description=_clean(values.get("description")),
        created_by=username,
    )
    return f"تم إنشاء المركز '{code}'"


def _create_project(command_bus: Any, uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.application.projects.commands import CreateProjectCommand

    code = _clean(values.get("code"))
    name = _clean(values.get("name"))
    if name is None or len(name) < 2:
        name = "مشروع " + (code or "جديد")
    if code:
        existing = None
        try:
            from core.domain.projects.value_objects import ProjectCode
            existing = uow.projects.get_by_code(ProjectCode(code))
        except Exception as e:
            logger.debug(f"project dup check failed: {e}")
        if existing:
            return f"تم تجاهل المشروع '{code}' (موجود مسبقاً)"
    budget = _parse_number(values.get("budget_amount"))
    base_currency = options.get("base_currency") or BASE_CURRENCY_FALLBACK
    bcc = _clean(values.get("budget_currency"))
    budget_currency = bcc if bcc and len(bcc) == 3 else base_currency
    status_map = {
        "نشط": "active", "قيد التنفيذ": "active", "قيد الانجاز": "active", "jاري العمل": "active",
        "مكتمل": "completed", "منتهي": "completed", "معلق": "on_hold",
        "ملغي": "cancelled", "تم الالغاء": "cancelled", "مؤرشف": "archived",
    }
    raw_status = _clean(values.get("status")) or "planning"
    status = status_map.get(raw_status.lower(), raw_status)
    valid_statuses = ("planning", "active", "on_hold", "completed", "cancelled", "archived")
    if status not in valid_statuses:
        raise ImportRowError(f"حالة مشروع غير صالحة: {raw_status}")
    tags = values.get("tags")
    tag_list = None
    if tags:
        tag_list = [t.strip() for t in re.split(r"[,;،]", tags) if t.strip()]
    _dispatch_command(
        command_bus, CreateProjectCommand,
        name=name, code=code,
        description=_clean(values.get("description")),
        customer_name=_clean(values.get("customer_name")),
        manager_name=_clean(values.get("manager_name")),
        budget_amount=Decimal(budget if budget is not None else 0),
        budget_currency=budget_currency,
        start_date=_parse_date(values.get("start_date")),
        end_date=_parse_date(values.get("end_date")),
        status=status, tags=tag_list,
        created_by=username,
    )
    return f"تم إنشاء المشروع '{name}'"


def _create_currency(command_bus: Any, uow: Any, values: Dict[str, str], username: str, options: Dict[str, Any]) -> str:
    from core.application.currency.commands import CreateCurrencyCommand

    name = _clean(values.get("name"))
    code = _clean(values.get("code"))
    if code is None or len(code) != 3 or not code.isalpha():
        derived = None
        if name:
            letters = re.sub(r"[^a-zA-Z]", "", name).upper()
            derived = letters[:3] if len(letters) >= 3 else None
        if derived is None:
            raise ImportRowError("كود العملة مطلوب (3 أحرف) ولم يُتمكن من توليده من الاسم")
        code = derived
    if name is None or len(name) < 2:
        name = code
    code = code.upper()
    existing = None
    try:
        from core.domain.currency.value_objects import CurrencyCode
        existing = uow.currencies.get_by_code(CurrencyCode(code))
    except Exception as e:
        logger.debug(f"currency dup check failed: {e}")
    if existing:
        return f"تم تجاهل العملة '{code}' (موجودة مسبقاً)"
    decimal_places = _parse_int(values.get("decimal_places"))
    is_base = _parse_bool(values.get("is_base"))
    _dispatch_command(
        command_bus, CreateCurrencyCommand,
        code=code, name=name, symbol=_clean(values.get("symbol")) or "",
        decimal_places=decimal_places if decimal_places is not None else 2,
        is_base=is_base if is_base is not None else False,
        created_by=username,
    )
    return f"تم إنشاء العملة '{code}'"


IMPORTERS: Dict[str, Callable[..., str]] = {
    "customers": _create_customer,
    "suppliers": _create_supplier,
    "products": _create_product,
    "accounts": _create_account,
    "sites": _create_site,
    "centers": _create_center,
    "projects": _create_project,
    "currencies": _create_currency,
}

COMMAND_ENTITIES = {"accounts", "sites", "centers", "projects", "currencies"}
UOW_ENTITIES = {"customers", "suppliers", "products"}


# =============================================================================
# النقاط النهائية
# =============================================================================


@router.post("/api/import/{entity}", response_model=ApiResponse)
async def import_entity(
    entity: str,
    file: UploadFile = File(...),
    mapping: str = Form("{}"),
    options: str = Form("{}"),
    current_user: dict = Depends(get_current_user),
    _auth: object = require_permission("system_config"),
):
    """يرفع ملف Excel ويستورد البيانات الرئيسية دفعة واحدة حسب الكيان."""
    if entity not in SUPPORTED_ENTITIES:
        return ApiResponse(success=False, message=f"كيان غير مدعوم: {entity}")

    try:
        import openpyxl

        mapping_data: Dict[str, int] = {}
        try:
            mapping_data = json.loads(mapping or "{}")
        except Exception:
            mapping_data = {}
        mapping_data = {str(k): int(v) for k, v in mapping_data.items() if v is not None}

        options_data: Dict[str, Any] = {}
        try:
            options_data = json.loads(options or "{}")
        except Exception:
            options_data = {}
        options_data["update_existing"] = bool(options_data.get("update_existing", True))

        contents = await file.read()
        if not contents:
            return ApiResponse(success=False, message="الملف فارغ")
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        if not wb.sheetnames:
            return ApiResponse(success=False, message="الملف لا يحتوي على أوراق عمل")
        ws = wb[wb.sheetnames[0]]
        raw_rows: List[List[str]] = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        if not raw_rows:
            return ApiResponse(success=False, message="الملف فارغ")

        # اكتشاف صف الرأس: أول صف فيه 2+ خلية نصية أو تطابق ترويسة معروفة.
        # عند وجود header_row في الخيارات (من معاينة الواجهة) يُعتمد عليه مباشرة.
        header_index = int(options_data.get("header_row", -1)) if options_data.get("header_row") is not None else -1
        if header_index < 0:
            header_index = 0
            for i, row in enumerate(raw_rows):
                non_empty = sum(1 for c in row if _clean(c))
                aliases = ENTITY_ALIASES.get(entity, {})
                matched = any(
                    _norm_header(c) in {_norm_header(a) for a in aliases[a]}
                    for c in row
                    for a in aliases
                )
                if non_empty >= 2 or (non_empty >= 1 and matched):
                    header_index = i
                    break
        header_index = max(0, min(header_index, len(raw_rows) - 1))
        while header_index > 0 and all(not _clean(c) for c in raw_rows[header_index]):
            header_index -= 1

        headers = [c.strip() for c in raw_rows[header_index]]
        mapping_data = _auto_map(headers, entity, mapping_data)

        data_rows = raw_rows[header_index + 1:]
        data_rows = [r for r in data_rows if any(_clean(c) for c in r)]

        username = current_user.get("username", "system")
        command_bus = bootstrap.container.resolve("command_bus") if entity in COMMAND_ENTITIES else None

        results: List[Dict[str, Any]] = []
        success = 0
        failed = 0
        start = datetime.now()

        for idx, row in enumerate(data_rows):
            # تعداد كبير الحجم: نحافظ على الاستمرارية عبر صفوف غير قابلة للفشل الشامل.
            try:
                values: Dict[str, str] = {}
                for field, col in mapping_data.items():
                    if 0 <= col < len(row):
                        values[field] = row[col]
                with bootstrap.uow() as uow:
                    if entity in UOW_ENTITIES:
                        message = IMPORTERS[entity](uow, values, username, options_data)
                    else:
                        message = IMPORTERS[entity](command_bus, uow, values, username, options_data)
                results.append({"row": header_index + idx + 2, "ok": True, "message": message})
                success += 1
            except ImportRowError as e:
                results.append({"row": header_index + idx + 2, "ok": False, "message": str(e)})
                failed += 1
            except Exception as e:
                logger.debug(f"import row {idx} failed: {e}")
                msg = str(e)
                if len(msg) > 300:
                    msg = msg[:300]
                results.append({"row": header_index + idx + 2, "ok": False, "message": msg})
                failed += 1

        return ApiResponse(
            success=True,
            message=f"اكتمل الاستيراد: {success} ناجح، {failed} مرفوض",
            data={
                "entity": entity,
                "total": len(data_rows),
                "success": success,
                "failed": failed,
                "durationMs": (datetime.now() - start).total_seconds() * 1000,
                "results": results,
            },
        )
    except Exception as e:
        logger.error(f"Import {entity} error: {e}", exc_info=True)
        return ApiResponse(success=False, message=str(e), errors=[str(e)])